# -*- coding: utf-8 -*-
"""
Lõi: Agent 06 - Xuất bản 1-Click (In-Memory Export Engine)
Định vị: .\agents\agent_6_export.py
Tiêu chuẩn: Zero-Disk I/O | Dynamic Merging | Fail-Fast | CV-5512
"""

import sqlite3
import json
import io
from pathlib import Path
from typing import Dict, Any

# Fail-Fast: Nếu thiếu thư viện, văng lỗi để Cổng Giáo Viên kích hoạt Khung Chẩn Đoán
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    raise ImportError("python-docx")

DB_PATH = Path(r".\yourdream.db")

class AgentExport:
    def __init__(self):
        self.db_path = DB_PATH

    def _get_connection(self):
        if not self.db_path.exists():
            raise FileNotFoundError(f"TỬ HUYỆT: Không tìm thấy CSDL Master tại {self.db_path}")
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _merge_data(self, teacher_id: str, item_id: str, cursor: sqlite3.Cursor) -> Dict[str, Any]:
        """Kéo dữ liệu Master và hòa trộn với Local Overrides trên RAM."""
        cursor.execute("SELECT * FROM atomic_nodes WHERE node_id = ?", (item_id,))
        base_row = cursor.fetchone()
        if not base_row:
            return {}
        
        final_data = dict(base_row)

        try:
            cursor.execute(
                "SELECT patch_data FROM local_overrides WHERE teacher_id = ? AND item_id = ?", 
                (teacher_id, item_id)
            )
            patch_row = cursor.fetchone()
            if patch_row and patch_row["patch_data"]:
                patch_dict = json.loads(patch_row["patch_data"])
                for key, val in patch_dict.items():
                    if val and str(val).strip(): 
                        final_data[key] = val
        except sqlite3.OperationalError:
            pass 

        return final_data

    def _build_cv5512_docx_bytes(self, data: Dict[str, Any]) -> bytes:
        """Đúc tệp Word chuẩn Công văn 5512 trên luồng RAM."""
        doc = Document()
        
        # Tiêu đề
        title = doc.add_heading("KẾ HOẠCH BÀI DẠY (CHUẨN CV 5512)", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph(f"Bài học/Chủ đề: {data.get('lesson_title', 'Không tên')}")
        doc.add_paragraph(f"Môn học: {data.get('subject', 'Toán')} - Khối lớp: {data.get('grade', 6)}")
        doc.add_paragraph(f"Mã Nút Tri Thức: {data.get('node_id', '')}")
        
        # Bảng Barem 4 Bước
        doc.add_heading("I. TIẾN TRÌNH DẠY HỌC & BAREM 4 BƯỚC", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Tiêu chí / Bước giải'
        hdr_cells[1].text = 'Nội dung chi tiết (0.25đ/bước)'
        
        steps = [
            ("Bước 1", data.get("barem_step_1", "")),
            ("Bước 2", data.get("barem_step_2", "")),
            ("Bước 3", data.get("barem_step_3", "")),
            ("Bước 4", data.get("barem_step_4", ""))
        ]
        
        for step_name, step_content in steps:
            if step_content:
                row_cells = table.add_row().cells
                row_cells[0].text = step_name
                row_cells[1].text = step_content

        # Bẫy ngộ nhận & Ghi chú Giáo viên
        doc.add_heading("II. LƯU Ý SƯ PHẠM & BẪY NGỘ NHẬN", level=2)
        doc.add_paragraph(f"Bẫy số 0 (Trap Zero): {data.get('trap_zero_description', 'Không có cảnh báo.')}")
        
        if data.get("teacher_note"):
            doc.add_heading("III. GHI CHÚ CÁ NHÂN (LOCAL OVERRIDE)", level=2)
            p = doc.add_paragraph(data.get("teacher_note"))
            p.runs[0].italic = True

        # Ghi ra luồng ByteIO
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def export_lesson_plan_bytes(self, teacher_id: str, item_id: str, format_type: str) -> Dict[str, Any]:
        """Đầu mối xuất bản 1-Click."""
        conn = self._get_connection()
        cursor = conn.cursor()
        
        try:
            final_data = self._merge_data(teacher_id, item_id, cursor)
            if not final_data:
                return {"status": "FAIL", "error": f"Không tìm thấy Nút Tri Thức {item_id}."}

            if format_type == "WORD_CV5512":
                file_bytes = self._build_cv5512_docx_bytes(final_data)
                mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                file_ext = "docx"
            elif format_type == "PDF_A4_PRINT":
                # TODO: Sẽ triển khai ReportLab ở phiên bản sau
                return {"status": "FAIL", "error": "Tính năng xuất PDF đang bảo trì."}
            else:
                return {"status": "FAIL", "error": "Định dạng không hỗ trợ."}
            
            clean_title = str(final_data.get('lesson_title', item_id)).replace(" ", "_").replace(":", "")
            file_name = f"CV5512_{clean_title}_{teacher_id}.{file_ext}"
            
            return {
                "status": "SUCCESS",
                "file_name": file_name,
                "mime_type": mime_type,
                "file_bytes": file_bytes
            }
        except Exception as e:
            return {"status": "FAIL", "error": str(e)}
        finally:
            conn.close()

agent_export = AgentExport()